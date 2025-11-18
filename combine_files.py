import os
import pandas as pd
from pathlib import Path

# Try to import openpyxl for Excel formatting
try:
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

def format_excel_file(filename):
    """
    Apply pretty formatting to an Excel file:
    - Bold headers with colored background
    - Auto-width columns
    - Frozen header row
    - Text wrapping for long cells
    - Borders
    """
    if not OPENPYXL_AVAILABLE:
        return
    
    try:
        # Load the workbook
        wb = load_workbook(filename)
        ws = wb.active
        
        # Define styles
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        wrap_alignment = Alignment(wrap_text=True, vertical='top')
        
        # Format header row
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        
        # Format data cells
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = border
                cell.alignment = wrap_alignment
                # Set font for better readability
                cell.font = Font(size=10)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    # Get the length of the cell value
                    if cell.value:
                        cell_length = len(str(cell.value))
                        # For wrapped text, consider a reasonable max width
                        if cell_length > 50:
                            cell_length = 50  # Cap at 50 characters for width
                        if cell_length > max_length:
                            max_length = cell_length
                except:
                    pass
            
            # Set column width with some padding
            # Minimum width of 10, maximum of 50
            adjusted_width = min(max(max_length + 2, 10), 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Freeze the header row
        ws.freeze_panes = 'A2'
        
        # Save the formatted workbook
        wb.save(filename)
        
    except Exception as e:
        print(f"  ⚠ Warning: Could not format Excel file: {e}")

def combine_folder_files(folder_path, output_folder):
    """
    Combine all CSV and XLSX files from a folder into single files.
    
    Args:
        folder_path: Path to the folder containing files to combine
        output_folder: Path to the output folder where combined files will be saved
    """
    folder_name = os.path.basename(folder_path)
    print(f"\n{'='*60}")
    print(f"Processing folder: {folder_name}")
    print(f"{'='*60}")
    
    # Get all CSV and XLSX files in the folder
    csv_files = sorted([f for f in os.listdir(folder_path) if f.endswith('.csv')])
    xlsx_files = sorted([f for f in os.listdir(folder_path) if f.endswith('.xlsx')])
    
    if not csv_files and not xlsx_files:
        print(f"  ⚠ No CSV or XLSX files found in '{folder_name}'")
        return
    
    # Combine CSV files
    if csv_files:
        print(f"  Found {len(csv_files)} CSV file(s)")
        combined_csv_data = []
        
        for csv_file in csv_files:
            csv_path = os.path.join(folder_path, csv_file)
            try:
                df = pd.read_csv(csv_path, encoding='utf-8-sig')
                combined_csv_data.append(df)
                print(f"    ✓ Loaded {len(df)} rows from {csv_file}")
            except Exception as e:
                print(f"    ⚠ Error reading {csv_file}: {e}")
        
        if combined_csv_data:
            # Combine all DataFrames
            combined_df = pd.concat(combined_csv_data, ignore_index=True)
            print(f"  Total rows combined: {len(combined_df)}")
            
            # Save combined CSV
            output_csv = os.path.join(output_folder, f"{folder_name}.csv")
            combined_df.to_csv(output_csv, index=False, encoding='utf-8-sig')
            print(f"  ✓ Combined CSV saved: '{output_csv}'")
    
    # Combine XLSX files
    if xlsx_files:
        print(f"  Found {len(xlsx_files)} XLSX file(s)")
        combined_xlsx_data = []
        
        for xlsx_file in xlsx_files:
            xlsx_path = os.path.join(folder_path, xlsx_file)
            try:
                df = pd.read_excel(xlsx_path, engine='openpyxl')
                combined_xlsx_data.append(df)
                print(f"    ✓ Loaded {len(df)} rows from {xlsx_file}")
            except Exception as e:
                print(f"    ⚠ Error reading {xlsx_file}: {e}")
        
        if combined_xlsx_data:
            # Combine all DataFrames
            combined_df = pd.concat(combined_xlsx_data, ignore_index=True)
            print(f"  Total rows combined: {len(combined_df)}")
            
            # Save combined XLSX
            output_xlsx = os.path.join(output_folder, f"{folder_name}.xlsx")
            combined_df.to_excel(output_xlsx, index=False, engine='openpyxl')
            
            # Apply formatting
            format_excel_file(output_xlsx)
            print(f"  ✓ Combined XLSX saved: '{output_xlsx}'")

def main():
    """
    Main function to combine files from all query folders.
    """
    # Get current directory
    current_dir = os.getcwd()
    print(f"Current directory: {current_dir}")
    
    # Define output folder
    output_folder = "combined_output"
    os.makedirs(output_folder, exist_ok=True)
    print(f"Output folder: {output_folder}")
    
    # Find all folders in current directory (excluding output folder and hidden/system folders)
    folders = []
    for item in os.listdir(current_dir):
        item_path = os.path.join(current_dir, item)
        if os.path.isdir(item_path) and item != output_folder and not item.startswith('.'):
            # Check if folder contains CSV or XLSX files
            has_data_files = any(f.endswith(('.csv', '.xlsx')) for f in os.listdir(item_path))
            if has_data_files:
                folders.append(item_path)
    
    if not folders:
        print("\n⚠ No folders with CSV/XLSX files found in current directory")
        return
    
    print(f"\nFound {len(folders)} folder(s) to process:")
    for folder in folders:
        print(f"  - {os.path.basename(folder)}")
    
    # Process each folder
    for folder_path in folders:
        combine_folder_files(folder_path, output_folder)
    
    print(f"\n{'='*60}")
    print("All folders processed!")
    print(f"Combined files saved to: '{output_folder}'")
    print(f"{'='*60}")

if __name__ == "__main__":
    main()

